# -*- coding: utf-8 -*-
"""
无界面自检模块（诊断用，不创建 QApplication）。

用途：验证运行环境（开发态或 PyInstaller 冻结版）中全部重型依赖可正常导入与
读写回环，重点覆盖函数内延迟导入、静态分析不可见的库（pdfplumber/pdfminer、
docx、reportlab、openpyxl、xlrd、xlsxwriter）。

入口：python main.py --smoke-test（冻结版：ChecklistTool.exe --smoke-test）
结果：明细写入 user_data/logs/smoke_test.log（末尾 SMOKE_TEST_OK / SMOKE_TEST_FAIL 标记）。
返回：0=全部通过，1=任一失败。GUI 版无控制台，验证以退出码 + 日志标记为准。
"""

import os
import sys
import traceback

from config import LOG_DIR, EXPORT_DIR, RESOURCE_DIR


def run_smoke_test() -> int:
    lines = []
    ok = True

    def step(name, fn):
        nonlocal ok
        try:
            fn()
            lines.append(f"[PASS] {name}")
        except Exception as e:
            ok = False
            lines.append(f"[FAIL] {name}: {e}")
            lines.append(traceback.format_exc())

    def ensure_module(mod: str):
        __import__(mod)

    # 1) 重型依赖导入（打包后静态分析不可见，逐项验证）
    for mod in (
        "PyQt6.QtWidgets",
        "pandas",
        "openpyxl",
        "xlrd",
        "xlsxwriter",
        "docx",
        "pdfplumber",
        "reportlab",
        "pdfminer.cmapdb",
    ):
        step(f"导入 {mod}", lambda m=mod: ensure_module(m))

    # 2) 随包资源（spec datas 是否正确打包）
    def check_icon():
        icon = os.path.join(RESOURCE_DIR, "assets", "icon.ico")
        if not os.path.isfile(icon):
            raise FileNotFoundError(f"缺少资源: {icon}")

    step("随包资源 assets/icon.ico", check_icon)

    # 3) xlsx 写入 + 解析回环（pandas + openpyxl + 项目解析器）
    def xlsx_roundtrip():
        import pandas as pd
        from core.parsers import load_table_from_file

        path = os.path.join(EXPORT_DIR, "smoke_test_input.xlsx")
        pd.DataFrame(
            {"系统": ["给排水", "电气"], "专业": ["消防", "照明"], "管径": [50, 25]}
        ).to_excel(path, index=False)
        df, cols, _ = load_table_from_file(path, header_rows=1)
        assert len(df) == 2, f"行数错误: {len(df)}"
        assert "系统" in cols, f"列缺失: {cols}"

    step("xlsx 写入+解析回环", xlsx_roundtrip)

    # 4) Excel 导出（xlsxwriter）
    def excel_export():
        import pandas as pd
        from core.export import export_to_excel

        df = pd.DataFrame({"系统": ["给排水"], "专业": ["消防"]})
        out = export_to_excel(df, "smoke_test_out.xlsx")
        assert os.path.isfile(out), f"导出文件不存在: {out}"

    step("Excel 导出", excel_export)

    # 5) PDF 导出 + pdfplumber 读回（reportlab + pdfplumber/pdfminer）
    def pdf_roundtrip():
        import pandas as pd
        import pdfplumber
        from core.export import export_to_pdf

        df = pd.DataFrame({"系统": ["给排水", "电气"], "专业": ["消防", "照明"]})
        out = export_to_pdf(df, "smoke_test_out.pdf")
        assert os.path.isfile(out), f"导出文件不存在: {out}"
        with pdfplumber.open(out) as pdf:
            assert len(pdf.pages) >= 1, "PDF 无页面"
            text = pdf.pages[0].extract_text() or ""
            assert text.strip(), "PDF 页面无文本"
            # 中文必须真实渲染（未注册 CJK 字体时 reportlab 会输出占位字符）
            assert "给排水" in text, f"PDF 中文渲染异常: {text[:60]!r}"

    step("PDF 导出+pdfplumber 读回", pdf_roundtrip)

    # 6) docx 建表 + 解析回环（python-docx + lxml）
    def docx_roundtrip():
        from docx import Document
        from core.parsers import load_table_from_file

        path = os.path.join(EXPORT_DIR, "smoke_test.docx")
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "系统"
        table.cell(0, 1).text = "专业"
        table.cell(1, 0).text = "给排水"
        table.cell(1, 1).text = "消防"
        table.cell(2, 0).text = "电气"
        table.cell(2, 1).text = "照明"
        doc.save(path)
        df, cols, _ = load_table_from_file(path)
        assert "系统" in cols, f"列缺失: {cols}"

    step("docx 建表+解析回环", docx_roundtrip)

    # 7) 设备数据 SQLite 存储层回环（建表/批量/分页/重命名/级联删除）
    def db_roundtrip():
        from core.data import (
            db_connect,
            init_db,
            create_list,
            get_list,
            list_lists,
            insert_rows_batch,
            get_rows_page,
            update_row,
            update_columns,
            rename_column,
            delete_list,
            set_template_columns,
            get_template_columns,
            rows_to_df,
            dedupe_columns,
            DataError,
        )

        db_path = os.path.join(EXPORT_DIR, "smoke_checklist.sqlite")
        for suffix in ("", "-wal", "-shm"):
            try:
                os.remove(db_path + suffix)
            except OSError:
                pass
        conn = db_connect(db_path)
        try:
            init_db(conn)
            # 列名去重
            assert dedupe_columns(["系统", "系统", "专业"]) == ["系统", "系统_2", "专业"]
            # 建清单
            lst = create_list(conn, "冒烟清单", ["系统", "专业"])
            assert lst["columns"] == ["系统", "专业"] and lst["row_count"] == 0
            # 重复列拒绝
            try:
                create_list(conn, "坏清单", ["A", "A"])
                raise AssertionError("重复列应被拒绝")
            except ValueError:
                pass
            # 批量插入 1000 行
            ids = insert_rows_batch(
                conn, lst["id"],
                [{"系统": f"系统{i % 7}", "专业": f"专业{i % 5}"} for i in range(1000)],
                editor_id="SMOKE",
            )
            assert len(ids) == 1000
            # 分页遍历（id 键翻页）
            collected, total, after = 0, 0, 0
            while True:
                page, total = get_rows_page(conn, lst["id"], after_id=after, limit=200)
                collected += len(page)
                if len(page) < 200:
                    break
                after = page[-1]["id"]
            assert collected == 1000 and total == 1000, f"{collected}/{total}"
            # 更新行
            first = get_rows_page(conn, lst["id"], limit=1)[0][0]
            update_row(conn, lst["id"], first["id"], {"系统": "给排水", "专业": "消防"})
            row = get_rows_page(conn, lst["id"], limit=1)[0][0]
            assert row["id"] == first["id"] and row["values"]["系统"] == "给排水"
            # 列重命名 re-key
            rename_column(conn, lst["id"], "系统", "系统名称")
            row = get_rows_page(conn, lst["id"], limit=1)[0][0]
            assert "系统名称" in row["values"] and "系统" not in row["values"]
            # 列结构替换
            update_columns(conn, lst["id"], ["系统名称", "专业", "备注"])
            # rows_to_df：index == 行 id
            rows = get_rows_page(conn, lst["id"], limit=10)[0]
            df = rows_to_df(rows, ["系统名称", "专业", "备注"])
            assert list(df.index) == [r["id"] for r in rows]
            assert list(df.columns) == ["系统名称", "专业", "备注"]
            # 模板列
            set_template_columns(conn, ["系统名称", "专业"])
            assert get_template_columns(conn) == ["系统名称", "专业"]
            # 不存在的清单
            try:
                update_columns(conn, 999999, ["A"])
                raise AssertionError("不存在的清单应报错")
            except DataError:
                pass
            # 级联删除
            assert delete_list(conn, lst["id"]) is True
            assert get_list(conn, lst["id"]) is None
            cascade = conn.execute(
                "SELECT COUNT(*) FROM list_rows WHERE list_id=?", (lst["id"],)
            ).fetchone()[0]
            assert cascade == 0
            assert list_lists(conn) == []
        finally:
            conn.close()
            for suffix in ("", "-wal", "-shm"):
                try:
                    os.remove(db_path + suffix)
                except OSError:
                    pass

    step("SQLite 设备数据存储层回环", db_roundtrip)

    lines.append("SMOKE_TEST_OK" if ok else "SMOKE_TEST_FAIL")
    log_text = "\n".join(lines) + "\n"

    log_path = os.path.join(LOG_DIR, "smoke_test.log")
    try:
        with open(log_path, "w", encoding="utf-8") as f:
            f.write(log_text)
    except Exception:
        pass

    # 开发态可见；冻结 GUI 版无控制台，以日志与退出码为准
    print(log_text)
    return 0 if ok else 1
